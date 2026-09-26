from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "DGFY-Dine-In-Workflow.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_borders(element):
    p_pr = element._element.get_or_add_pPr() if hasattr(element, "_element") else element
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_run_font(run, name="Arial", size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=4, line_spacing=1.05):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def add_text(paragraph, text, bold=False, color="000000", size=10.5):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(paragraph, space_after=5, line_spacing=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True, color="000000")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_step(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    style_paragraph(paragraph, space_after=2.5, line_spacing=1.03)
    add_text(paragraph, f"{number}. ", bold=True, size=10.3)
    add_text(paragraph, text, size=10.3)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, space_after=2.5, line_spacing=1.03)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    add_text(paragraph, text, size=10.3)
    return paragraph


def add_payment_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.55), Inches(2.35), Inches(2.95)]
    headers = ["Payment option", "When available", "Order result"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, cell in enumerate(header.cells):
        cell.width = widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(p, space_after=0, line_spacing=1.0)
        add_text(p, headers[idx], bold=True, color="FFFFFF", size=9.4)

    rows = [
        ("Online full", "Always, if enabled", "Paid after provider success."),
        ("Online downpayment", "If enabled", "Partial payment; balance due at service."),
        ("Pay at store", "Only when Support In-Store Payment is on", "Unpaid until cashier records payment."),
    ]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].width = widths[idx]
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cells[idx])
            set_cell_margins(cells[idx], top=95, bottom=95)
            if row_index % 2 == 1:
                set_cell_shading(cells[idx], "F3F6F9")
            p = cells[idx].paragraphs[0]
            style_paragraph(p, space_after=0, line_spacing=1.0)
            add_text(p, value, size=9.3)
    return table


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(paragraph, space_after=0, line_spacing=1.0)
    add_text(paragraph, "DGFY | Dine-in workflow", color="666666", size=8.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
add_footer(section)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Heading 1", "Heading 2", "Title"):
    styles[style_name].font.name = "Arial"
    styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
title.paragraph_format.keep_with_next = True
remove_borders(title)
remove_borders(doc.styles["Title"])
add_text(title, "DGFY Dine In Order Flow", bold=True, size=21, color="000000")

subtitle = doc.add_paragraph()
style_paragraph(subtitle, space_after=10, line_spacing=1.05)
add_text(subtitle, "A short guide for branch setup, customer checkout, POS preparation, and payment handling.", color="404040", size=10.5)

intro = doc.add_paragraph()
style_paragraph(intro, space_after=8, line_spacing=1.08)
add_text(intro, "Core rule: ", bold=True, size=10.5)
add_text(intro, "Dine-in controls where and when the order is served. Payment policy controls how and when it is paid.", size=10.5)

add_heading(doc, "1 Activate Dine In", 1)
for number, text in [
    (1, "Open Settings, then Locations or Branches, and select the branch."),
    (2, "Turn on Supports Dine-in and Allow Scheduled Orders."),
    (3, "Choose whether Immediate Fulfillment is on. Turn it off for scheduled-only Dine-in."),
    (4, "Set branch hours, preparation or lead time, timezone, and order cutoff rules."),
    (5, "Under Payment Policies, configure Online Full, Online Downpayment, and Support In-Store Payment."),
    (6, "Save and publish. Confirm the branch shows Dine-in on the storefront."),
]:
    add_step(doc, number, text)

add_heading(doc, "2 Storefront Customer Flow", 1)
for number, text in [
    (1, "Open dgfy.ph, choose the business, then choose the exact branch."),
    (2, "Select items, quantities, and modifiers, then open the cart."),
    (3, "At checkout, choose Dine-in."),
    (4, "Choose Now or a future serve-by date and time within the branch rules."),
    (5, "Enter contact details and any order notes."),
    (6, "Choose an available payment option from the table below."),
    (7, "Review the branch, serve-by time, items, discounts, VAT, total, and balance, then place the order."),
    (8, "Show the customer the order PIN, serve-by time, payment status, and remaining balance."),
]:
    add_step(doc, number, text)

add_payment_table(doc)

add_heading(doc, "3 POS Cashier Flow", 1)
for number, text in [
    (1, "Sign in, open the terminal shift, and confirm the terminal is online."),
    (2, "Open Orders, then Dine-in Schedule."),
    (3, "Work from the queue sorted by serve-by time. Each order shows an H/M/S countdown and its payment status."),
    (4, "Use the modal reminders at 15, 5, 3, and 1 minute before service."),
    (5, "Open the order, verify the PIN, branch, items, notes, table, serve-by time, and balance."),
    (6, "Start preparation before the serve-by time and mark the order Ready to Serve when complete."),
    (7, "When the customer arrives, confirm or assign the table, collect any balance, record the actual payment method, and provide the receipt."),
    (8, "Mark the order Served or Completed only after service is finished."),
]:
    add_step(doc, number, text)

doc.add_page_break()
add_heading(doc, "4 Status Rules", 1)
add_bullet(doc, "Fulfillment: Placed → Confirmed → Preparing → Ready to Serve → Served.")
add_bullet(doc, "Payment: Unpaid → Partially Paid → Paid. Payment can also be Pending, Failed, or Refunded.")
add_bullet(doc, "Do not mark an order paid just because it was placed. Pay-at-store orders remain unpaid until the cashier records payment.")
add_bullet(doc, "The storefront should normally not require a table number before arrival. The POS can assign or confirm the table.")

add_heading(doc, "5 Example", 1)
example = doc.add_paragraph()
style_paragraph(example, space_after=2, line_spacing=1.05)
add_text(example, "A customer selects Dine-in for 7:00 PM. The POS alerts the cashier at 6:45, 6:55, 6:57, and 6:59. The cashier prepares the order before 7:00, marks it Ready to Serve, collects any outstanding amount when the customer arrives, and marks it Served after handover.", size=10.3)

doc.core_properties.title = "DGFY Dine In Order Flow"
doc.core_properties.subject = "Dine-in branch activation, storefront checkout, POS operations, and payment rules"
doc.core_properties.author = "DGFY"
doc.core_properties.keywords = "DGFY, dine-in, POS, storefront, payments"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
