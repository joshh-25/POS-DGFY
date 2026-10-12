from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
PROJECT = "POS Customer Display Synchronization"
VERSION = "1.0"
DATE = "2026-10-07"
REPO = r"C:\Users\USER\.codex\worktrees\updates-pos\POS-DGFY"
ACCENT = "1A4E8D"
LIGHT_BLUE = "EAF2FB"
LIGHT_GRAY = "F5F7FA"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, space_before, space_after in (
        ("Title", 22, 0, 10),
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.text = f"{PROJECT} | Plan v{VERSION} | {DATE}"
    for run in footer_p.runs:
        set_font(run, size=8, color="64748B")


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_font(run, size=22, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    run2 = p2.add_run(subtitle)
    set_font(run2, size=10.5, color="475569")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_font(r)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_font(r)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, ACCENT)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(label))
        set_font(r, size=9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=9.2)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_metadata(doc, section_number, title, purpose, dependencies):
    add_table(
        doc,
        ["Plan field", "Value"],
        [
            ("Project Name", PROJECT),
            ("Plan Version", VERSION),
            ("Section Number", section_number),
            ("Section Title", title),
            ("Purpose", purpose),
            ("Dependencies", dependencies),
            ("Repository", REPO),
            ("Baseline date", DATE),
        ],
        widths=[1.55, 5.55],
    )


def add_required_sections(doc, confirmed, assumptions, requirements, components, risks, tests, acceptance, execution, crossrefs):
    add_heading(doc, "Confirmed Facts", 1)
    add_bullets(doc, confirmed)
    add_heading(doc, "Assumptions and Unknowns", 1)
    add_bullets(doc, assumptions)
    add_heading(doc, "Requirements Covered", 1)
    add_table(doc, ["ID", "Requirement"], requirements, widths=[1.05, 6.05])
    add_heading(doc, "Affected System Components", 1)
    add_table(doc, ["Component", "Impact"], components, widths=[2.25, 4.85])
    add_heading(doc, "Risks and Mitigations", 1)
    add_table(doc, ["Risk", "Mitigation"], risks, widths=[2.45, 4.65])
    add_heading(doc, "Tests Required", 1)
    add_bullets(doc, tests)
    add_heading(doc, "Acceptance Criteria", 1)
    add_table(doc, ["ID", "Measurable result"], acceptance, widths=[1.05, 6.05])
    add_heading(doc, "Execution Dependencies", 1)
    add_bullets(doc, execution)
    add_heading(doc, "Cross References", 1)
    add_bullets(doc, crossrefs)


def make_doc(filename, section_number, title, purpose, dependencies, confirmed, assumptions, requirements, body, components, risks, tests, acceptance, execution, crossrefs):
    doc = Document()
    format_document(doc)
    add_title(doc, title, f"{PROJECT} | Section {section_number} | Plan v{VERSION}")
    add_metadata(doc, section_number, title, purpose, dependencies)
    add_heading(doc, "Plan Content", 1)
    body(doc)
    add_required_sections(doc, confirmed, assumptions, requirements, components, risks, tests, acceptance, execution, crossrefs)
    doc.save(ROOT / filename)


common_confirmed = [
    "The active repository/worktree for this plan is C:\\Users\\USER\\.codex\\worktrees\\updates-pos\\POS-DGFY.",
    "The POS app is owned by apps/dgfy-pos, while shared POS business and UI code is under packages/web-core/src/features/pos.",
    "apps/dgfy-pos/src/main.jsx currently mounts a HashRouter with cashier routes at / and /terminal; no customer-display route exists.",
    "TerminalPage owns the POS shell and composes TerminalOperationsWorkspace and the checkout terminal.",
    "usePosTerminalViewModel owns the active cart, cart totals, checkout state, and checkout lifecycle used by the cashier UI.",
    "POSCheckoutTerminalView renders the current sale item list, line totals, subtotal, discounts, VAT breakdown, and total.",
    "The current code has a tenant-scoped POS catalog SSE stream at /pos/catalog/events, but no cart or customer-display synchronization channel.",
    "The backend remains the source of truth for sales, inventory, transactions, terminal shifts, and transaction history.",
    "No source code was changed while this plan was prepared.",
]

common_assumptions = [
    "MVP means a second browser window or tab on the same POS workstation and same registered terminal, normally shown on a second physical monitor.",
    "The customer display is read-only and does not accept customer input, payment, checkout commands, or cashier actions.",
    "The display should show the active in-store cashier sale before checkout, not the public Storefront tracking page.",
    "The display should show item names, quantities, line totals, subtotal, applicable discounts, VAT summary, and the current total.",
    "The customer display should not show customer phone, email, address, internal notes, cashier credentials, access tokens, or payment secrets.",
    "Separate-device LAN or remote customer displays are future scope unless the product owner explicitly expands the requirement.",
]

common_requirements = [
    ("FR-001", "An authorized cashier can open a dedicated customer-display window from the active POS terminal."),
    ("FR-002", "The display mirrors the cashier's active cart within a bounded real-time latency target of one second under normal same-workstation conditions."),
    ("FR-003", "The display shows each current item, quantity, unit price, line total, subtotal, discounts where applicable, and the current total."),
    ("FR-004", "Cart add, quantity change, item removal, discount change, clear-sale, and checkout-state changes update the display without a manual refresh."),
    ("FR-005", "A newly opened display receives the current full snapshot through a ready handshake and does not wait for the next cart mutation."),
    ("FR-006", "The display provides explicit waiting, empty-cart, active-sale, payment-in-progress, completed, disconnected, and session-ended states."),
    ("BR-001", "The cashier cart and its existing derived totals remain the source for the display projection; the display does not create a competing sale or recalculate authoritative money."),
    ("BR-002", "Checkout, payment, inventory deduction, receipt issuance, and transaction persistence remain unchanged and cashier-controlled."),
    ("SEC-001", "No DGFY access token, refresh token, pairing cookie, password, or sensitive customer identity is broadcast through the display channel."),
    ("SEC-002", "A display session is scoped to one tenant and terminal session and ignores snapshots for another session or stale revision."),
    ("UI-001", "The customer display is optimized for distance viewing, full-screen use, keyboard-free operation, high contrast, and responsive desktop/large-monitor layouts."),
    ("NFR-001", "The sync protocol is bounded, versioned, idempotent, and resilient to display reloads, cashier reloads, popup blocking, offline transitions, and channel reconnection."),
]

common_components = [
    ("apps/dgfy-pos/src/main.jsx", "Add the customer-display route and preserve the existing cashier routes."),
    ("packages/web-core/src/features/pos/pages/TerminalPage.jsx", "Expose the open-display action and coordinate display-session lifecycle."),
    ("packages/web-core/src/features/pos/hooks/usePosTerminalViewModel.js", "Provide the existing cart and derived-total snapshot to the sync adapter."),
    ("packages/web-core/src/features/pos/components/POSCheckoutTerminalView.jsx", "Keep cashier cart behavior intact; add only the customer-display launch/status affordance where the existing shell owns it."),
    ("packages/web-core/src/features/pos/services", "Add the local display channel/session adapter and snapshot validation."),
    ("packages/web-core/src/features/pos/utils", "Add a pure customer-display projection that reuses existing cart/financial display helpers."),
    ("apps/dgfy-api", "No MVP mutation/API/schema change; validate that checkout and transaction source-of-truth contracts remain untouched."),
    ("tests/frontend-cross-app and POS tests", "Add unit, component, and browser evidence for the two-window flow and negative security cases."),
]

common_risks = [
    ("Stale or out-of-order snapshots", "Include a session id and monotonic revision; ignore mismatched or older snapshots and resync on ready."),
    ("Cross-tenant data leakage", "Generate an unpredictable in-memory display session per cashier terminal; never use a global fixed channel; clear on tenant/session change."),
    ("Popup blocked or customer window closed", "Show an actionable cashier status and keep checkout usable; allow reopening without changing the sale."),
    ("Financial display drift", "Project from the existing view-model totals and test discounts, VAT, decimal quantities, modifiers, and empty carts."),
    ("Sensitive data on a public-facing screen", "Allowlist display fields; omit identity, notes, payment details, tokens, and internal operational metadata."),
    ("Older iMin/browser runtime", "Use a small transport adapter with BroadcastChannel as the primary path and a same-origin fallback only if supported by the browser baseline; verify on the POS runtime."),
    ("Display remains stale after lock/logout", "Send session-ended/clear messages and make the display return to a neutral waiting state when heartbeats or session ownership disappear."),
]

common_tests = [
    "Unit tests for the projection allowlist, money/quantity formatting, revision ordering, session scoping, and clear-state behavior.",
    "Component tests for waiting, empty cart, active cart, payment, completed, disconnected, and session-ended display states.",
    "Cashier component tests for open-display, popup-blocked, display-reconnected, and display-closed behavior.",
    "Two-page Playwright test using one cashier page and one customer-display page on the same POS origin.",
    "Negative security test proving customer display input cannot mutate the cashier cart or call checkout actions.",
    "Lifecycle test covering add item, quantity update, remove item, discount change, clear sale, checkout, and new sale.",
    "Responsive screenshot/DOM evidence at desktop and one constrained POS/mobile-sized viewport, plus console/page-error checks.",
    "POS production build and existing POS/web-core regression suites; run architecture checks for any backend or governed architecture files.",
]

common_acceptance = [
    ("AC-001", "Cashier opens the customer-display window and the customer screen shows a waiting or empty state without exposing cashier controls."),
    ("AC-002", "Adding two different items in the cashier cart causes the same item names, quantities, line totals, and current total to appear on the customer display within one second."),
    ("AC-003", "Changing quantity, removing an item, applying a discount, and clearing the sale produces the corresponding customer-display state without a page refresh."),
    ("AC-004", "A display opened after the cashier already has items receives the full current snapshot immediately after the ready handshake."),
    ("AC-005", "Cashier checkout remains the only path that creates the transaction; the customer display has no checkout, payment, edit, or navigation mutation."),
    ("AC-006", "Closing, locking, logging out, changing company, or losing the display session clears the display to a neutral state and never leaves the previous tenant's cart visible."),
    ("AC-007", "Popup failure, channel interruption, page reload, and stale messages are visible as recoverable display states and do not block cashier checkout."),
    ("AC-008", "The target POS and shared web-core builds/tests pass, with no new architecture-boundary or security exceptions."),
]

common_execution = [
    "Phase 476 remains in_progress in the authoritative ledger; the plan's next eligible continuous phase is Phase 477 only after Phase 476 eligibility is reconciled.",
    "No database migration or production data backfill is expected for the same-workstation MVP.",
    "If the product owner requires a separate LAN device, stop and extend the design with backend-mediated pairing and transport before implementation.",
    "Implementation must preserve the POS terminal pairing, shift, location, permission, and compliance rules in ADRs 0031, 0044, and 0007.",
]

common_crossrefs = [
    "docs/START_HERE.md",
    "docs/architecture/ARCHITECTURE_BOUNDARIES.md",
    "docs/architecture/ARCHITECTURE_GOVERNANCE.md",
    "docs/architecture/adr/0001-modular-monolith-boundaries.md",
    "docs/architecture/adr/0025-pos-application-shells-and-lan-host-runtime.md",
    "docs/architecture/adr/0031-pos-terminal-pairing-and-shift-safe-navigation.md",
    "docs/architecture/adr/0044-pos-terminal-device-pairing.md",
    "docs/architecture/adr/0055-tenant-scoped-pos-catalog-realtime-invalidation.md",
    "docs/features/IMPLEMENTATION_PHASE_LEDGER.md",
    "packages/web-core/src/features/pos/components/POSCheckoutTerminalView.jsx",
    "packages/web-core/src/features/pos/hooks/usePosTerminalViewModel.js",
]


def body_current(doc):
    add_heading(doc, "Critical Assessment", 2)
    add_paragraph(doc, "The earlier live verification proved Storefront tracking and POS online-order synchronization. It did not prove a two-screen in-store POS workflow. Repository and runtime inspection confirm that the requested cashier/customer-display surface is not present today.")
    add_paragraph(doc, "The existing POS already has the correct state to reuse: the cashier cart and its derived totals are held by usePosTerminalViewModel and rendered by POSCheckoutTerminalView. The missing capability is a read-only projection and transport to a second display, not a new order or payment model.")
    add_heading(doc, "Current Flow", 2)
    add_numbered(doc, [
        "apps/dgfy-pos/src/main.jsx mounts the POS HashRouter and sends / and /terminal to TerminalPage.",
        "TerminalPage establishes terminal/session/shift context and composes the operational POS workspace.",
        "TerminalOperationsWorkspace owns the checkout workspace composition and operational modes.",
        "POSCheckoutTerminal renders POSCheckoutTerminalView with the view model.",
        "usePosTerminalViewModel owns the active cart, line mutations, checkout totals, payment state, and checkout lifecycle.",
        "POSCheckoutTerminalView renders the cashier's Current Sale panel, line totals, financial breakdown, and Total.",
        "The existing checkout submit path uses createPosCheckout in packages/web-core/src/features/pos/services/posService.js and the backend POS checkout contract; this plan does not alter that path.",
    ])
    add_heading(doc, "Confirmed Gap", 2)
    add_bullets(doc, [
        "No /customer-display route or equivalent POS display page exists in apps/dgfy-pos/src/main.jsx.",
        "No customer-display component, second-screen launch action, display-session store, BroadcastChannel, postMessage bridge, or cart display event exists in the POS surface.",
        "The only matching Monitor Display Screen text in TerminalPage.jsx is inside a loading illustration and is not a feature.",
        "The existing /pos/catalog/events SSE stream is for catalog invalidation only and intentionally carries no cart or sale data.",
    ])
    add_heading(doc, "Target Objective", 2)
    add_paragraph(doc, "Add a read-only customer-facing POS display for the same registered terminal. A cashier opens it in a second window or tab, normally placed on a second monitor. The display mirrors the existing active cart and totals in real time while preserving the backend as the source of truth for checkout, payment, inventory, and transaction history.")
    add_heading(doc, "What Must Remain Unchanged", 2)
    add_bullets(doc, [
        "Cashier authentication, terminal selection, pairing, location grants, permissions, shift opening, checkout, payment, inventory deduction, receipt printing, audit, and compliance behavior.",
        "The public Storefront catalog and customer tracking routes.",
        "Existing transaction schema and backend checkout contracts for the same-workstation MVP.",
    ])


def body_requirements(doc):
    add_heading(doc, "Scope Decision", 2)
    add_paragraph(doc, "The recommended first release is local and same-terminal. It solves the actual two-screen cashier/customer scenario without introducing a second sale owner, a new database record, or an unauthenticated LAN protocol. A separate device over the network is explicitly deferred until its pairing and trust model are specified.")
    add_heading(doc, "In Scope", 2)
    add_bullets(doc, [
        "A cashier control to open or reopen the customer display.",
        "A display-only POS route with full-screen, distance-readable presentation.",
        "Real-time mirroring of the active cart and the existing cashier-derived financial totals.",
        "Clear lifecycle behavior for empty cart, active sale, payment progress, completion, clear sale, lock, logout, company switch, reload, and disconnect.",
        "Accessibility, responsive layout, popup failure handling, and local browser diagnostics.",
    ])
    add_heading(doc, "Out of Scope", 2)
    add_bullets(doc, [
        "Customer touch input, self-service ordering, payment, tipping, receipt reprint, or cashier actions.",
        "Cross-device or remote LAN synchronization, cloud display registration, or display administration.",
        "Storefront tracking, online order status, delivery tracking, or post-checkout customer accounts.",
        "Persisting display snapshots or adding a transaction/display database table.",
    ])
    add_heading(doc, "Functional and Business Requirements", 2)
    add_table(doc, ["ID", "Requirement", "Priority"], [
        ("FR-001 to FR-006", "Open, initialize, mirror, update, and recover the customer display.", "P1"),
        ("BR-001 to BR-002", "Reuse cashier cart/totals; keep checkout and transaction authority unchanged.", "P0"),
        ("SEC-001 to SEC-002", "Do not broadcast credentials or sensitive data; scope and order sessions.", "P0"),
        ("UI-001", "Provide distance-readable, read-only customer presentation.", "P1"),
        ("NFR-001", "Use a versioned, bounded, idempotent sync protocol.", "P1"),
    ], widths=[1.3, 4.8, 0.9])
    add_heading(doc, "Hidden Requirements Made Explicit", 2)
    add_bullets(doc, [
        "The customer display must not show stale data from a previous tenant, terminal, cashier, or sale.",
        "Opening a display must not clear, pause, or alter the cashier cart.",
        "A display must recover from reload without waiting for a new cashier action.",
        "The customer must be able to read the final amount at a distance without seeing internal tax/accounting fields that are not intended for customer display.",
        "The cashier must be able to continue checkout when the display window is blocked, closed, offline, or unsupported.",
    ])


def body_architecture(doc):
    add_heading(doc, "Recommended Design", 2)
    add_paragraph(doc, "Use a same-origin local display session. The cashier creates an unpredictable display-session identifier in memory, opens the existing POS app in a second window at a dedicated customer-display route, and sends read-only snapshots over a small transport adapter. BroadcastChannel is the primary transport for modern same-origin windows. A narrowly scoped same-origin fallback may be added only if the supported iMin/browser baseline requires it; no credentials or payment data may be placed in fallback storage.")
    add_heading(doc, "Proposed Modules and Files", 2)
    add_table(doc, ["Area", "Planned change"], [
        ("POS route", "Add a display-only route in apps/dgfy-pos/src/main.jsx, likely /customer-display, without changing cashier route ownership."),
        ("Display page", "Add a POS-owned page/component for waiting, empty, active, payment, completed, stale, and disconnected states."),
        ("Sync adapter", "Add packages/web-core/src/features/pos/services/customerDisplaySync.js for session creation, handshake, snapshot send/receive, revision checks, and teardown."),
        ("Projection", "Add packages/web-core/src/features/pos/utils/customerDisplayProjection.js to allowlist and normalize display-safe fields from the existing view model."),
        ("Cashier integration", "Wire the launch/status control through TerminalPage and the existing checkout shell; do not duplicate cart state."),
        ("Tests", "Add web-core unit/component contracts and POS Playwright evidence for two pages and lifecycle transitions."),
        ("Architecture record", "Create a new ADR after confirming no existing ADR covers this local customer-display transport decision. The next ADR index number is 0090."),
    ], widths=[1.35, 5.65])
    add_heading(doc, "Sync Message Contract", 2)
    add_table(doc, ["Message", "Direction", "Purpose"], [
        ("display.ready", "Display to cashier", "Proves the opened display is listening and requests the current full snapshot."),
        ("display.snapshot", "Cashier to display", "Carries the complete read-only cart projection and state."),
        ("display.clear", "Cashier to display", "Clears the previous sale on clear, lock, logout, tenant switch, or session end."),
        ("display.close", "Either side", "Signals teardown; cashier remains fully usable if the display disappears."),
        ("display.heartbeat", "Both sides", "Optional liveness signal for a bounded stale/disconnected state."),
    ], widths=[1.45, 1.55, 4.0])
    add_heading(doc, "Snapshot Shape", 2)
    add_bullets(doc, [
        "protocol_version and schema_version for forward compatibility.",
        "display_session_id and terminal_id for session scoping; do not include access tokens or pairing cookies.",
        "state: waiting, empty, active, payment_pending, completed, disconnected, or session_ended.",
        "revision and updated_at for monotonic ordering and diagnostics.",
        "store_name and currency as display context.",
        "items: stable line key, customer-safe item name, quantity, unit price, line total, and approved modifier text if needed.",
        "subtotal, allowed discount summary, allowed VAT presentation, and total from the existing view-model financial projection.",
        "No customer phone, email, address, internal note, cashier identity, card/payment secret, auth token, or backend-only audit field.",
    ])
    add_heading(doc, "State Ownership", 2)
    add_table(doc, ["State", "Owner", "Customer display behavior"], [
        ("Active cart", "Cashier view model", "Render the latest accepted snapshot."),
        ("Financial total", "Existing POS calculation/checkout contract", "Render the existing derived total; never create a second money calculation."),
        ("Checkout/payment", "Cashier POS", "Show payment-in-progress or a neutral processing state; no customer action."),
        ("Completed sale", "Cashier checkout result", "Show a short completion confirmation and final total, then clear on next sale."),
        ("Session/tenant/terminal", "POS auth and terminal lifecycle", "Clear immediately when ownership becomes invalid."),
    ], widths=[1.25, 2.25, 3.5])
    add_heading(doc, "Why No MVP Database Change", 2)
    add_paragraph(doc, "The display is a transient presentation of an active cart, not a business record. Persisting it would create a competing source of truth and introduce cleanup, tenant scoping, and historical meaning that the requested workflow does not need. The backend remains authoritative when checkout is submitted.")


def body_frontend(doc):
    add_heading(doc, "Cashier Flow", 2)
    add_numbered(doc, [
        "Cashier unlocks the existing POS terminal and opens a shift under the existing terminal, location, permission, and compliance rules.",
        "Cashier selects Open Customer Display from the checkout shell.",
        "POS creates a session id in memory and opens /customer-display with the session id in the route state.",
        "The customer display sends display.ready; the cashier sends the complete current snapshot.",
        "The cashier continues using the existing catalog/cart controls. Every meaningful cart/view-model state change publishes a new projection.",
        "The cashier can reopen the display if the popup is blocked or the window is closed. Reopening never clears or changes the cart.",
        "At clear sale, lock, logout, company switch, or session replacement, the cashier sends display.clear/session_ended.",
    ])
    add_heading(doc, "Customer Display States", 2)
    add_table(doc, ["State", "Visible content", "Required behavior"], [
        ("Waiting", "Store branding and a simple waiting message.", "No item or prior-sale data; safe before handshake."),
        ("Empty", "Store branding and 'Ready for your order' or equivalent.", "No controls that can mutate the sale."),
        ("Active", "Item rows, quantities, line totals, subtotal, discounts if applicable, and prominent total.", "Update without refresh; use large typography and high contrast."),
        ("Payment", "Processing message and current total.", "Do not show sensitive payment details or allow interaction."),
        ("Completed", "Thank-you/complete message and final total.", "Clear when the cashier begins the next sale."),
        ("Disconnected", "Display connection warning and safe last-known/neutral state.", "Never pretend an old snapshot is current; allow automatic resync."),
        ("Session ended", "Neutral waiting state.", "Remove all previous tenant/sale details."),
    ], widths=[1.0, 2.65, 3.35])
    add_heading(doc, "Distance-Readable UX", 2)
    add_bullets(doc, [
        "Use full-screen layout with no cashier sidebar, settings, navigation, or browser-dependent controls.",
        "Make the current total the strongest visual element after the item list; keep item names and quantities readable at a distance.",
        "Use responsive layout for a wide secondary monitor and a constrained tablet-sized test viewport.",
        "Use semantic headings and regions, visible focus handling for the cashier launch control, and reduced-motion-safe transitions.",
        "Do not expose debug identifiers, raw event payloads, customer PII, or internal error stacks.",
    ])
    add_heading(doc, "Integration Boundaries", 2)
    add_bullets(doc, [
        "The projection adapter consumes the existing view-model values rather than reimplementing cart totals.",
        "The display page receives sanitized snapshots and renders presentation only.",
        "The existing checkout controls, payment dialogs, receipt renderer, incoming queue, and Storefront tracking remain unchanged.",
        "The route must be able to mount without a cashier cart and must not block the main cashier route if its chunk or transport fails.",
    ])


def body_security_tests(doc):
    add_heading(doc, "Security and Authorization", 2)
    add_bullets(doc, [
        "Opening the display is an authorized POS-terminal action; use the already resolved terminal/shift/session context rather than introducing a second credential authority.",
        "The display route is read-only. It must not expose or reuse cashier action handlers, checkout mutation functions, or payment controls.",
        "Use a cryptographically unpredictable, short-lived, in-memory display-session identifier. Do not use a tenant id alone as a channel name.",
        "Do not place bearer tokens, refresh tokens, pairing cookies, passwords, or secrets in BroadcastChannel payloads, query strings, localStorage, or sessionStorage.",
        "Validate every inbound display message against the current display-session identifier and expected message type. Ignore malformed, stale, or foreign messages.",
        "Allowlist item names and modifier labels as text. Render them as text nodes; do not inject HTML or trust client-supplied markup.",
        "Clear all display state on tenant/company change, terminal lock, logout, session expiry, popup teardown, and display-session replacement.",
        "Do not show customer contact information, delivery address, internal notes, cashier identity, internal audit fields, or payment-provider details.",
    ])
    add_heading(doc, "Verification Matrix", 2)
    add_table(doc, ["Requirement", "Evidence"], [
        ("FR-001, FR-005", "POS component test and two-page Playwright launch/ready handshake."),
        ("FR-002 to FR-004", "Two-page Playwright lifecycle test with bounded latency assertion and DOM snapshots."),
        ("FR-006", "Display-state component tests and reload/disconnect Playwright scenarios."),
        ("BR-001, BR-002", "Projection contract plus regression tests proving createPosCheckout and existing checkout controls remain unchanged."),
        ("SEC-001, SEC-002", "Static payload allowlist test, foreign-session rejection test, tenant/terminal lifecycle test."),
        ("UI-001", "Desktop and constrained viewport screenshots plus accessibility/semantic DOM assertions."),
        ("NFR-001", "Revision ordering, duplicate message, reconnect, popup-blocked, and browser-support tests."),
    ], widths=[1.5, 5.5])
    add_heading(doc, "GIVEN WHEN THEN Scenarios", 2)
    add_bullets(doc, [
        "GIVEN a cashier has an empty sale, WHEN the customer display opens, THEN it shows the neutral/empty state and not a previous order.",
        "GIVEN the cashier adds item A and item B, WHEN the snapshot arrives, THEN the customer display shows both items, their quantities, line totals, subtotal, and current total.",
        "GIVEN the cashier changes quantity or removes a line, WHEN a newer revision arrives, THEN the display replaces the prior projection without duplicate rows.",
        "GIVEN a stale revision arrives after a newer revision, WHEN the display receives it, THEN it ignores the stale revision.",
        "GIVEN the cashier changes company or locks the terminal, WHEN the lifecycle event occurs, THEN the customer display clears all prior data.",
        "GIVEN the display window is blocked or closed, WHEN the cashier attempts to open it, THEN checkout remains available and the cashier receives a recoverable status.",
        "GIVEN the display sends an edit or checkout-shaped message, WHEN the cashier receives it, THEN it ignores the message and the cart/checkout remain unchanged.",
    ])
    add_heading(doc, "Build and Runtime Proof", 2)
    add_bullets(doc, [
        "Run the shared web-core/POS unit and component suites covering changed files.",
        "Build apps/dgfy-pos and any other app surface that consumes changed web-core code.",
        "Run the local stack with cashier and customer-display pages on the same POS origin.",
        "Collect screenshots, DOM assertions, console/page-error output, latency measurements, and a clean-worktree check.",
        "Run npm run check:architecture for architecture-sensitive changes; run docs lint if ADR or governed docs are changed.",
    ])


def body_phases(doc):
    add_heading(doc, "Continuous Phase Numbering", 2)
    add_paragraph(doc, "The authoritative ledger currently reaches Phase 476, which is still marked in_progress. The next continuous number is Phase 477. This plan should be recorded as the candidate for Phase 477 only after Phase 476's remaining acceptance gates are reconciled; it must not reset numbering or mark Phase 477 completed prematurely.")
    add_heading(doc, "Implementation Phases", 2)
    add_table(doc, ["Phase", "Objective", "Exit criteria"], [
        ("PHASE-477-A", "Architecture and contract", "Scope is confirmed as same-terminal; ADR 0090 is drafted if the new local transport decision is accepted; snapshot/message contract is reviewed."),
        ("PHASE-477-B", "Projection and transport foundation", "Pure allowlisted projection, session id, handshake, revision ordering, teardown, and fallback adapter have unit coverage."),
        ("PHASE-477-C", "Customer-display route and UI", "Display route renders all states responsively and never exposes cashier mutations or sensitive fields."),
        ("PHASE-477-D", "Cashier launch and lifecycle integration", "Cashier can open/reopen/observe the display; cart and checkout lifecycle publishes correct snapshots; failures do not block checkout."),
        ("PHASE-477-E", "Hardening and verification", "Security negatives, two-page E2E, responsive evidence, builds, architecture checks, and regression suites pass."),
        ("PHASE-477-F", "Local rollout and rollback proof", "Local pilot is completed on a two-monitor or two-window setup; operator instructions and rollback/disable behavior are recorded."),
    ], widths=[1.3, 2.35, 3.35])
    add_heading(doc, "Dependency Graph", 2)
    add_paragraph(doc, "Architecture and snapshot contract -> projection and transport -> display route -> cashier integration -> browser/security verification -> local pilot. Unit projection tests can run in parallel with the display visual shell. The two-page E2E and runtime pilot depend on the route, transport, and cashier lifecycle wiring.")
    add_heading(doc, "Rollout", 2)
    add_bullets(doc, [
        "Keep the feature disabled until the customer-display route and cashier launch control are both present and verified.",
        "Pilot on one local terminal with two browser windows/monitors while cashier checkout remains the existing path.",
        "Monitor popup failures, transport disconnects, stale revision counts, page errors, and customer-display session teardown events.",
        "Do not require a schema migration or production data backfill for the local MVP.",
    ])
    add_heading(doc, "Rollback", 2)
    add_bullets(doc, [
        "Disable or remove the cashier launch control and route while leaving the existing cashier checkout untouched.",
        "If the transport fails, the cashier continues without a customer display; no transaction rollback is needed because the display is presentation-only.",
        "If stale or cross-tenant data is observed, immediately stop the customer-display route, clear browser sessions, and revert the display integration before resuming pilot use.",
    ])
    add_heading(doc, "Open Product Decision", 2)
    add_paragraph(doc, "Before implementation, confirm whether 'second screen' means a second window on the same Windows POS workstation or a separate device connected over LAN. This plan selects the same-workstation interpretation because it matches the stated cashier-plus-customer monitor workflow and minimizes security and operational risk. A separate-device decision changes the transport, pairing, authentication, network failure, and rollout design.")


def body_index(doc):
    add_heading(doc, "Original User Request", 2)
    add_paragraph(doc, "Create a plan for a real POS customer-facing display: the POS has two screens, one for the cashier and one for the customer, and the customer screen must stay synchronized with the cashier's active sale.")
    add_heading(doc, "Final Objective", 2)
    add_paragraph(doc, "Plan the addition of a read-only, same-terminal customer-display mode that mirrors the cashier's current POS cart, item quantities, line totals, and final total in real time without creating a second source of truth or changing checkout authority.")
    add_heading(doc, "Baseline and Decision Summary", 2)
    add_table(doc, ["Decision", "Plan position"], [
        ("Current state", "No dedicated POS customer-display route or synchronization feature exists."),
        ("Recommended MVP", "Same-origin second window/tab on the same registered terminal."),
        ("Transport", "In-memory session plus BroadcastChannel adapter, with only a justified same-origin fallback."),
        ("Source of truth", "Existing cashier view-model cart/totals; backend remains authoritative at checkout."),
        ("Persistence", "No new database table, migration, backfill, or historical record."),
        ("Security", "Read-only display, session-scoped messages, no credential or sensitive identity broadcast."),
        ("Next phase", "Candidate Phase 477, after the ledger's Phase 476 in_progress eligibility is reconciled."),
    ], widths=[1.55, 5.55])
    add_heading(doc, "Reading Order", 2)
    add_numbered(doc, [
        "01 Objective and Current State",
        "02 Requirements and Business Rules",
        "03 Target Architecture and Sync Protocol",
        "04 Frontend UX and Integration",
        "05 Security Testing and Verification",
        "06 Implementation Phases Rollout and Acceptance",
    ])
    add_heading(doc, "Requirement Index", 2)
    add_table(doc, ["Group", "Coverage"], [
        ("FR-001 to FR-006", "Opening, handshake, mirroring, lifecycle, and recovery."),
        ("BR-001 to BR-002", "Cart/totals source of truth and unchanged checkout ownership."),
        ("SEC-001 to SEC-002", "No secrets/sensitive data and session/tenant isolation."),
        ("UI-001", "Distance-readable, responsive, read-only customer presentation."),
        ("NFR-001", "Versioned, bounded, idempotent, reconnect-safe synchronization."),
    ], widths=[1.7, 5.4])
    add_heading(doc, "ADR and Governance Index", 2)
    add_bullets(doc, [
        "ADR 0001 governs routes -> controllers -> usecases -> repositories -> models for any backend work.",
        "ADR 0025 keeps backend ownership of sales/inventory/transactions and clients responsible for UI/checkout workflow.",
        "ADR 0031 and ADR 0044 govern terminal identity, pairing, shift, location, permission, and security boundaries.",
        "ADR 0055 is a relevant realtime precedent for scoped, reconnectable POS event delivery, but its catalog invalidation payload must not be reused for cart data.",
        "A new ADR 0090 is recommended if the local customer-display transport becomes a durable architecture decision.",
    ])
    add_heading(doc, "Warnings", 2)
    add_bullets(doc, [
        "This is a plan-only artifact. No application code, tests, database schema, or phase-ledger entry was changed.",
        "The plan intentionally does not claim that a customer display exists today.",
        "If the requested second screen is a separate LAN device rather than a second window on the same workstation, implementation must stop at Phase 477-A and expand the trust/transport design.",
    ])


def build_all():
    make_doc(
        "00_POS_Customer_Display_Synchronization_Index.docx",
        "00",
        "POS Customer Display Synchronization Plan",
        "Master index for an execution-ready plan to add a synchronized cashier/customer POS display.",
        "Read this index first, then the six numbered sections in order.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_index,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "01_Objective_and_Current_State.docx",
        "01",
        "Objective and Current State",
        "Establish the confirmed system baseline, current cashier flow, missing capability, and target objective.",
        "Repository inspection and the POS runtime at port 5174.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_current,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "02_Requirements_and_Business_Rules.docx",
        "02",
        "Requirements and Business Rules",
        "Turn the requested two-screen behavior into testable functional, business, security, UI, and non-functional requirements.",
        "Section 01 baseline; product confirmation of the same-terminal MVP assumption.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_requirements,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "03_Target_Architecture_and_Sync_Protocol.docx",
        "03",
        "Target Architecture and Sync Protocol",
        "Define the recommended local transport, snapshot contract, state ownership, and module/file boundaries.",
        "Sections 01 and 02; ADR 0001, ADR 0025, ADR 0031, ADR 0044, and ADR 0055.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_architecture,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "04_Frontend_UX_and_Integration.docx",
        "04",
        "Frontend UX and Integration",
        "Describe the cashier launch flow, customer-display states, distance-readable UX, and integration boundaries.",
        "Sections 01 to 03; existing POS terminal shell and view-model contracts.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_frontend,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "05_Security_Testing_and_Verification.docx",
        "05",
        "Security Testing and Verification",
        "Define security boundaries, failure handling, verification evidence, and browser test scenarios.",
        "Sections 01 to 04; existing terminal pairing, permissions, and architecture gates.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_security_tests,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )
    make_doc(
        "06_Implementation_Phases_Rollout_and_Acceptance.docx",
        "06",
        "Implementation Phases Rollout and Acceptance",
        "Sequence the implementation, preserve continuous phase numbering, and define rollout, rollback, and final acceptance.",
        "Sections 01 to 05; Phase 476 ledger eligibility reconciliation.",
        common_confirmed,
        common_assumptions,
        common_requirements,
        body_phases,
        common_components,
        common_risks,
        common_tests,
        common_acceptance,
        common_execution,
        common_crossrefs,
    )


if __name__ == "__main__":
    build_all()
